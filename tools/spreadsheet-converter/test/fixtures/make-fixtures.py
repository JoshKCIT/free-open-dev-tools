#!/usr/bin/env python
"""Writes the .xlsx fixtures of the spreadsheet converter's tests, and re-reads any .xlsx file as a second opinion.

Two independent writers make the fixtures, so the reader is tested against files this folder's own writer never made:

  openpyxl 3.1.5   openpyxl-people.xlsx (sheet People, a hidden sheet and a very hidden sheet)
                   openpyxl-date1904.xlsx (the 1904 date system)
  XlsxWriter 3.x   xlsxwriter-sales.xlsx (shared strings, a rich string, formulas with cached values, an error,
                   a percent format, a dd/mm/yyyy date, a date with a time, serials 59, 60 and 61 as dates)

Usage, from a virtual environment that holds only these two packages (never the machine's own Python):

  python -m venv venv-13
  venv-13/Scripts/python -m pip install openpyxl==3.1.5 XlsxWriter
  venv-13/Scripts/python make-fixtures.py write          # writes the three files beside this script
  venv-13/Scripts/python make-fixtures.py check FILE     # opens FILE with openpyxl, prints every cell as JSON

The values written below are the expected values of the tests: they are literals typed here, what each writer was
asked to store, never read back from this folder's reader. `check` is how the writer test's expected values for the
committed tool-output.xlsx were obtained.
"""
import datetime
import json
import pathlib
import sys

HERE = pathlib.Path(__file__).resolve().parent


def write_people() -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append(["name", "count", "when", "flag", "ratio", "code", "note"])
    ws.append(["Ada", 36, datetime.date(2020, 1, 31), True, 98.5, "00123"])
    ws.append(
        [
            'Bo, "Q"',
            None,
            datetime.datetime(2021, 6, 1, 13, 45, 30),
            False,
            0.3333333333333333,
            "1e3",
            "line1\nline2",
        ]
    )
    ws.append(["Cy", 9007199254740992, datetime.time(8, 30, 0), True, -0.0, "  spaced  ", "é€漢"])
    hidden = wb.create_sheet("Hidden data")
    hidden.append(["secret", 42])
    hidden.sheet_state = "hidden"
    very = wb.create_sheet("Very hidden")
    very.append(["deep", 7])
    very.sheet_state = "veryHidden"
    wb.save(HERE / "openpyxl-people.xlsx")


def write_date1904() -> None:
    import openpyxl
    from openpyxl.utils.datetime import CALENDAR_MAC_1904

    wb = openpyxl.Workbook()
    wb.epoch = CALENDAR_MAC_1904
    ws = wb.active
    ws.title = "Mac dates"
    ws.append(["day", "epoch day", "afternoon"])
    ws.append([datetime.date(2020, 1, 31), datetime.date(1904, 1, 1), datetime.datetime(2000, 2, 29, 12, 0, 0)])
    wb.save(HERE / "openpyxl-date1904.xlsx")


def write_sales() -> None:
    import xlsxwriter

    wb = xlsxwriter.Workbook(str(HERE / "xlsxwriter-sales.xlsx"))
    ws = wb.add_worksheet("Sales")
    date_fmt = wb.add_format({"num_format": "dd/mm/yyyy"})
    stamp_fmt = wb.add_format({"num_format": "yyyy-mm-dd hh:mm:ss"})
    time_fmt = wb.add_format({"num_format": "hh:mm:ss"})
    pct_fmt = wb.add_format({"num_format": "0.0%"})
    bold = wb.add_format({"bold": True})
    italic = wb.add_format({"italic": True})

    for col, title in enumerate(["Item", "Date", "Stamp", "Rate", "Flag", "Sum", "Error", "Rich", "Text", "Plain", "Clock"]):
        ws.write(0, col, title)

    ws.write(1, 0, "Widget")
    ws.write_datetime(1, 1, datetime.datetime(2024, 2, 29), date_fmt)
    ws.write_datetime(1, 2, datetime.datetime(2024, 1, 1, 8, 30, 0), stamp_fmt)
    ws.write_number(1, 3, 0.125, pct_fmt)
    ws.write_boolean(1, 4, True)
    ws.write_formula(1, 5, "=1+2", None, 3)
    ws.write_formula(1, 6, "=1/0", None, "#DIV/0!")
    ws.write_rich_string(1, 7, bold, "Bold", " plain ", italic, "more")
    ws.write_string(1, 8, "ünï")
    ws.write_string(1, 9, "007")
    ws.write_datetime(1, 10, datetime.time(8, 30, 0), time_fmt)

    ws.write(2, 0, "Widget")
    ws.write_formula(2, 5, '=CONCATENATE("a","b")', None, "ab")
    ws.write_formula(2, 4, "=1=1", None, True)
    ws.write_string(2, 8, "<b>&amp;</b>")

    ws.write(3, 0, "Serials")
    ws.write_number(3, 1, 59, date_fmt)
    ws.write_number(3, 2, 60, date_fmt)
    ws.write_number(3, 3, 61, date_fmt)
    wb.close()


def check(path: str) -> None:
    import openpyxl

    sys.stdout.reconfigure(encoding="utf-8")  # the console code page cannot print every character

    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        cells = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                cells.append(
                    {
                        "ref": cell.coordinate,
                        "type": cell.data_type,
                        "value": cell.value,
                        "format": cell.number_format,
                    }
                )
        out.append({"sheet": ws.title, "state": ws.sheet_state, "cells": cells})
    print(json.dumps(out, ensure_ascii=False, default=str, indent=1))


def versions() -> None:
    import openpyxl

    try:
        import xlsxwriter

        writer = xlsxwriter.__version__
    except ImportError:
        writer = "not installed"
    print(f"Python {sys.version.split()[0]}, openpyxl {openpyxl.__version__}, XlsxWriter {writer}")


def main(argv: list) -> int:
    if len(argv) >= 2 and argv[1] == "write":
        write_people()
        write_date1904()
        write_sales()
        versions()
        return 0
    if len(argv) == 3 and argv[1] == "check":
        check(argv[2])
        return 0
    if len(argv) == 2 and argv[1] == "versions":
        versions()
        return 0
    print(__doc__)
    return 2


if __name__ == "__main__":
    sys.exit(main(sys.argv))
